"""Tabular export/import helpers (CSV + XLSX) shared by factory & admin panels."""

from __future__ import annotations

import csv
import io

MEDIA_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def write_table(headers: list[str], rows: list[list], fmt: str, sheet_name: str = "Sheet1") -> bytes:
    """Serialize a header + rows into CSV or XLSX bytes."""
    fmt = fmt.lower()
    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(headers)
        for row in rows:
            writer.writerow(["" if c is None else c for c in row])
        # utf-8-sig so Excel opens UTF-8 correctly (Uzbek/Cyrillic).
        return buf.getvalue().encode("utf-8-sig")

    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name[:31]
        ws.append(headers)
        head_fill = PatternFill("solid", fgColor="2563EB")
        head_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = head_fill
            cell.font = head_font
        for row in rows:
            ws.append(["" if c is None else c for c in row])
        # Reasonable column widths.
        for col_idx, header in enumerate(headers, start=1):
            width = max(12, min(40, len(str(header)) + 4))
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width
        ws.freeze_panes = "A2"
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    raise ValueError(f"Unsupported format: {fmt}")


def read_table(data: bytes, filename: str) -> list[list[str]]:
    """Parse an uploaded CSV or XLSX into rows of strings (header included)."""
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        rows: list[list[str]] = []
        for r in ws.iter_rows(values_only=True):
            if r is None:
                continue
            rows.append(["" if c is None else str(c).strip() for c in r])
        return rows

    # Default: CSV. Tolerate utf-8-sig BOM.
    text = data.decode("utf-8-sig", errors="replace")
    return [list(r) for r in csv.reader(io.StringIO(text))]
